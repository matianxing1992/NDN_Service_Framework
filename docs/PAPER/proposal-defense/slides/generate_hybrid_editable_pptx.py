#!/usr/bin/env python3
"""Generate a PDF-matched PPTX with all extractable text kept editable.

The PDF graphics are rendered as a full-slide background after Ghostscript has
removed every PDF text object with ``-dFILTERTEXT``.  Text is reconstructed
from ``pdftohtml -xml`` output, including source coordinates, font size, font
family, RGB color, bold, and italic styling. Adjacent visual lines that belong
to one paragraph share one wrapping PowerPoint textbox; styled fragments stay
as separate runs inside that textbox. Centered blocks retain paragraph
alignment, while explicit numbered items remain separate paragraphs with a
hanging indent.

Text baked into a raster image in the source PDF cannot be extracted and stays
in that image.  All PDF text objects, including diagram labels, table cells,
references, titles, and footers, are rebuilt as native PowerPoint content.
"""

from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import dataclass
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from add_speaker_notes_to_pptx import inject_notes, load_notes_from_tex


SLIDES_DIR = Path(__file__).resolve().parent
PDF_PATH = SLIDES_DIR / "main.pdf"
OUTPUT_PATH = SLIDES_DIR / "NDNSF_proposal_hybrid_editable.pptx"
BUILD_DIR = SLIDES_DIR / "build" / "hybrid_editable_pptx"
SLIDE_W_IN = 16.0
SLIDE_H_IN = 9.0
RENDER_DPI = "300"
PDFTOHTML_ZOOM = "1.5"
MANIFEST_NAME = "editable-text-manifest.json"
BUILD_MARKER_NAME = ".ndnsf-hybrid-pptx-build"
STANDARD_FOOTER_TOP_IN = 8.55
FOOTER_BOTTOM_MARGIN_IN = 0.08


@dataclass(frozen=True)
class FontSpec:
    source_id: str
    size: float
    source_family: str
    mapped_family: str
    rgb: tuple[int, int, int]
    family_bold: bool
    family_italic: bool


@dataclass(frozen=True)
class InlineTextRun:
    text: str
    bold: bool
    italic: bool


@dataclass(frozen=True)
class TextSpan:
    span_id: int
    page: int
    page_w: float
    page_h: float
    left: float
    top: float
    width: float
    height: float
    font: FontSpec
    inline_runs: tuple[InlineTextRun, ...]

    @property
    def right(self) -> float:
        return self.left + self.width

    @property
    def bottom(self) -> float:
        return self.top + self.height

    @property
    def text(self) -> str:
        return "".join(run.text for run in self.inline_runs)


@dataclass(frozen=True)
class PdfPage:
    number: int
    width: float
    height: float
    spans: tuple[TextSpan, ...]


@dataclass(frozen=True)
class PdfDocument:
    pages: tuple[PdfPage, ...]

    @property
    def spans(self) -> tuple[TextSpan, ...]:
        return tuple(span for page in self.pages for span in page.spans)


@dataclass(frozen=True)
class TextGroup:
    group_id: int
    page: int
    spans: tuple[TextSpan, ...]
    line_start_span_ids: tuple[int, ...] = ()
    alignment: str = "left"


@dataclass(frozen=True)
class BulletMarker:
    page: int
    span: TextSpan


@dataclass(frozen=True)
class ConversionPlan:
    groups_by_page: dict[int, tuple[TextGroup, ...]]
    bullets_by_page: dict[int, tuple[BulletMarker, ...]]
    assignments: dict[int, str]
    adjacent_pairs: int
    multi_span_groups: int
    svs_citation_groups: int


@dataclass(frozen=True)
class BuildStats:
    textboxes: int
    text_runs: int
    bullet_shapes: int
    background_images: int
    visual_region_images: int


@dataclass(frozen=True)
class VisualRegion:
    name: str
    left: float
    top: float
    width: float
    height: float


@dataclass(frozen=True)
class SlideVisualRegionSpec:
    required_text: tuple[str, ...]
    regions: tuple[VisualRegion, ...]


@dataclass(frozen=True)
class PreparedVisualRegion:
    region: VisualRegion
    image_path: Path
    left: float
    top: float
    width: float
    height: float


# These regions contain graphics that remain raster after PDF text filtering.
# Match on stable slide content instead of ordinal page numbers so the same map
# is safe for the full, shortened, and reordered proposal decks.
SLIDE_VISUAL_REGION_SPECS = (
    SlideVisualRegionSpec(
        required_text=("Four-Provider Mobility: Experiment Design",),
        regions=(
            VisualRegion(
                "VISUAL_REGION_MOBILITY_TOPOLOGY", 2.25, 2.05, 3.85, 4.10
            ),
        ),
    ),
    SlideVisualRegionSpec(
        required_text=("Mobility Results: Coverage and Switching",),
        regions=(
            VisualRegion(
                "VISUAL_REGION_MOBILITY_RESULTS", 2.60, 1.55, 11.75, 4.75
            ),
        ),
    ),
    SlideVisualRegionSpec(
        required_text=("Request Message", "After policy bootstrap"),
        regions=(
            VisualRegion(
                "VISUAL_REGION_REQUEST_PACKET", 9.58, 3.23, 4.54, 5.04
            ),
        ),
    ),
    SlideVisualRegionSpec(
        required_text=("Selective ACK", "Providers do not have to ACK"),
        regions=(
            VisualRegion(
                "VISUAL_REGION_ACK_PACKET", 9.50, 2.95, 4.55, 5.24
            ),
        ),
    ),
    SlideVisualRegionSpec(
        required_text=(
            "Custom Provider Selection Strategy",
            "Provider selection is user-side",
        ),
        regions=(
            VisualRegion(
                "VISUAL_REGION_SELECTION_PACKET", 9.50, 3.30, 4.55, 4.75
            ),
        ),
    ),
    SlideVisualRegionSpec(
        required_text=("Provider Selection Strategies: Built-In Choices",),
        regions=(
            VisualRegion(
                "VISUAL_REGION_FIRST_RESPONDING_SEQUENCE",
                1.56,
                2.23,
                6.31,
                3.76,
            ),
            VisualRegion(
                "VISUAL_REGION_RANDOM_SELECTION_SEQUENCE",
                8.15,
                2.27,
                6.28,
                3.76,
            ),
        ),
    ),
    SlideVisualRegionSpec(
        required_text=(
            "Provider Selection Strategies: Multi-Provider and Custom",
        ),
        regions=(
            VisualRegion(
                "VISUAL_REGION_ALL_RESPONDERS_SEQUENCE",
                1.58,
                2.25,
                6.24,
                3.73,
            ),
            VisualRegion(
                "VISUAL_REGION_CUSTOM_SELECTION_SEQUENCE",
                8.16,
                2.29,
                6.26,
                3.72,
            ),
        ),
    ),
    SlideVisualRegionSpec(
        required_text=("Network Loss: NDNSF v0.1 and Baselines",),
        regions=(
            VisualRegion(
                "VISUAL_REGION_NETWORK_TOPOLOGY", 1.37, 2.05, 5.26, 4.16
            ),
            VisualRegion(
                "VISUAL_REGION_NETWORK_SUCCESS_CHART",
                7.23,
                2.74,
                7.84,
                3.69,
            ),
        ),
    ),
    SlideVisualRegionSpec(
        required_text=("NDNSF v0.1 Evaluation: Historical Mobility Diagnostic",),
        regions=(
            VisualRegion(
                "VISUAL_REGION_MOBILITY_TOPOLOGY", 1.25, 2.18, 5.79, 2.86
            ),
            VisualRegion(
                "VISUAL_REGION_MOBILITY_RESULTS_TABLE",
                7.72,
                1.81,
                7.20,
                3.13,
            ),
        ),
    ),
    SlideVisualRegionSpec(
        required_text=("UAV Workload", "UAV-APP is used to expose"),
        regions=(
            VisualRegion(
                "VISUAL_REGION_UAV_WORKLOAD_DIAGRAM",
                1.96,
                1.38,
                12.08,
                5.13,
            ),
        ),
    ),
    SlideVisualRegionSpec(
        required_text=(
            "v0.1 Evidence: Network and Provider Failures",
            "Contextual network-loss comparison",
        ),
        regions=(
            VisualRegion(
                "VISUAL_REGION_SHORT_NETWORK_ROUTE",
                1.18,
                1.92,
                6.23,
                0.84,
            ),
            VisualRegion(
                "VISUAL_REGION_SHORT_NETWORK_RESULTS_TABLE_RULES",
                1.90,
                4.93,
                4.70,
                1.93,
            ),
            VisualRegion(
                "VISUAL_REGION_SHORT_PROVIDER_FAILURE_TABLE_RULES",
                9.00,
                3.39,
                4.95,
                1.90,
            ),
        ),
    ),
    SlideVisualRegionSpec(
        required_text=(
            "v0.1 Evidence: Mobility and Availability",
            "Provider availability",
        ),
        regions=(
            VisualRegion(
                "VISUAL_REGION_SHORT_MOBILITY_DIAGRAM",
                1.98,
                1.84,
                4.88,
                2.74,
            ),
            VisualRegion(
                "VISUAL_REGION_SHORT_MOBILITY_TABLE_RULES",
                2.10,
                5.06,
                4.65,
                2.00,
            ),
            VisualRegion(
                "VISUAL_REGION_SHORT_AVAILABILITY_TIMELINE",
                9.92,
                3.80,
                4.15,
                1.28,
            ),
            VisualRegion(
                "VISUAL_REGION_SHORT_AVAILABILITY_TABLE_RULES",
                9.47,
                5.20,
                4.20,
                2.03,
            ),
        ),
    ),
    SlideVisualRegionSpec(
        required_text=(
            "UAV Workload and Requirements",
            "Mobility and changing reachability",
        ),
        regions=(
            VisualRegion(
                "VISUAL_REGION_SHORT_UAV_COLLABORATION_DIAGRAM",
                1.08,
                2.30,
                7.40,
                3.95,
            ),
        ),
    ),
)


def local_tag(element: ET.Element) -> str:
    return element.tag.rsplit("}", 1)[-1].lower()


def strip_subset_prefix(family: str) -> str:
    return re.sub(r"^[A-Z]{6}\+", "", family)


def map_font_family(source_family: str) -> str:
    """Map TeX PDF font names to their corresponding installable families."""

    family = strip_subset_prefix(source_family)
    lowered = family.lower()
    if "lmsans" in lowered:
        return "Latin Modern Sans"
    if "lmmono" in lowered or lowered.startswith("cmtt"):
        return "Latin Modern Mono"
    if "lmroman" in lowered or re.match(r"^cmr\d", lowered):
        return "Latin Modern Roman"
    if "lmmath" in lowered or "msam" in lowered:
        return "Cambria Math"
    return re.sub(
        r"-(?:regular|bold|italic|oblique|bolditalic|boldoblique)$",
        "",
        family,
        flags=re.IGNORECASE,
    )


def parse_rgb(value: str) -> tuple[int, int, int]:
    value = value.strip()
    if not re.fullmatch(r"#[0-9a-fA-F]{6}", value):
        raise ValueError(f"Unsupported pdftohtml color: {value!r}")
    return tuple(int(value[index:index + 2], 16) for index in (1, 3, 5))


def parse_font_spec(element: ET.Element) -> FontSpec:
    source_family = element.attrib["family"]
    normalized = strip_subset_prefix(source_family).lower()
    return FontSpec(
        source_id=element.attrib["id"],
        size=float(element.attrib["size"]),
        source_family=source_family,
        mapped_family=map_font_family(source_family),
        rgb=parse_rgb(element.attrib.get("color", "#000000")),
        family_bold=bool(re.search(r"bold|demi|semibold", normalized)),
        family_italic=bool(re.search(r"italic|oblique|slanted", normalized)),
    )


def append_inline_run(
    runs: list[InlineTextRun], text: str | None, *, bold: bool, italic: bool
) -> None:
    if not text:
        return
    if runs and runs[-1].bold == bold and runs[-1].italic == italic:
        previous = runs[-1]
        runs[-1] = InlineTextRun(previous.text + text, bold, italic)
    else:
        runs.append(InlineTextRun(text, bold, italic))


def extract_inline_runs(element: ET.Element) -> tuple[InlineTextRun, ...]:
    runs: list[InlineTextRun] = []

    def walk(node: ET.Element, bold: bool, italic: bool) -> None:
        tag = local_tag(node)
        node_bold = bold or tag in {"b", "strong"}
        node_italic = italic or tag in {"i", "em"}
        append_inline_run(
            runs, node.text, bold=node_bold, italic=node_italic
        )
        for child in node:
            walk(child, node_bold, node_italic)
            append_inline_run(
                runs, child.tail, bold=node_bold, italic=node_italic
            )

    walk(element, False, False)
    return tuple(runs)


def is_strict_descendant(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
    except ValueError:
        return False
    return path != root


def reset_build_dir(protected_paths: tuple[Path, ...]) -> None:
    resolved = BUILD_DIR.resolve()
    repository_build_root = (SLIDES_DIR / "build").resolve()
    temporary_root = Path(tempfile.gettempdir()).resolve()
    inside_repository_build = is_strict_descendant(
        resolved, repository_build_root
    )
    inside_dedicated_temp = (
        is_strict_descendant(resolved, temporary_root)
        and resolved.name.startswith("ndnsf-")
    )
    if not (inside_repository_build or inside_dedicated_temp):
        raise ValueError(
            "Refusing to clear build directory outside the approved roots "
            f"{repository_build_root} and {temporary_root}/ndnsf-*: {resolved}"
        )

    for protected_path in protected_paths:
        protected = protected_path.resolve()
        if protected == resolved or is_strict_descendant(protected, resolved):
            raise ValueError(
                f"Build directory {resolved} overlaps protected path {protected}"
            )

    if BUILD_DIR.exists():
        marker = BUILD_DIR / BUILD_MARKER_NAME
        if inside_dedicated_temp and not marker.is_file():
            raise ValueError(
                "Refusing to clear an existing temporary directory without "
                f"the NDNSF build marker: {resolved}"
            )
        shutil.rmtree(BUILD_DIR)
    BUILD_DIR.mkdir(parents=True)
    (BUILD_DIR / BUILD_MARKER_NAME).write_text(
        "Created by generate_hybrid_editable_pptx.py; contents are disposable.\n",
        encoding="utf-8",
    )


def require_external_tools() -> None:
    missing = [
        executable
        for executable in ("gs", "pdftohtml", "pdftotext")
        if shutil.which(executable) is None
    ]
    if missing:
        raise RuntimeError(
            "Missing required executable(s): " + ", ".join(missing)
        )


def extract_styled_text() -> PdfDocument:
    xml_path = BUILD_DIR / "pdftohtml.xml"
    subprocess.run(
        [
            "pdftohtml",
            "-q",
            "-xml",
            "-noroundcoord",
            "-hidden",
            "-nomerge",
            "-i",
            "-fontfullname",
            "-enc",
            "UTF-8",
            "-zoom",
            PDFTOHTML_ZOOM,
            str(PDF_PATH),
            str(xml_path),
        ],
        check=True,
    )
    if not xml_path.exists():
        raise RuntimeError(f"pdftohtml did not create {xml_path}")

    root = ET.parse(xml_path).getroot()
    fonts: dict[str, FontSpec] = {}
    for element in root.iter():
        if local_tag(element) != "fontspec":
            continue
        parsed = parse_font_spec(element)
        previous = fonts.get(parsed.source_id)
        if previous is not None and previous != parsed:
            raise RuntimeError(
                f"pdftohtml reused font id {parsed.source_id} inconsistently"
            )
        fonts[parsed.source_id] = parsed

    pages: list[PdfPage] = []
    next_span_id = 1
    for page_element in root:
        if local_tag(page_element) != "page":
            continue
        page_number = int(page_element.attrib["number"])
        page_w = float(page_element.attrib["width"])
        page_h = float(page_element.attrib["height"])
        spans: list[TextSpan] = []
        for text_element in page_element.iter():
            if local_tag(text_element) != "text":
                continue
            inline_runs = extract_inline_runs(text_element)
            if not "".join(run.text for run in inline_runs).strip():
                continue
            font_id = text_element.attrib["font"]
            if font_id not in fonts:
                raise RuntimeError(
                    f"Page {page_number} text references unknown font {font_id}"
                )
            spans.append(
                TextSpan(
                    span_id=next_span_id,
                    page=page_number,
                    page_w=page_w,
                    page_h=page_h,
                    left=float(text_element.attrib["left"]),
                    top=float(text_element.attrib["top"]),
                    width=float(text_element.attrib["width"]),
                    height=float(text_element.attrib["height"]),
                    font=fonts[font_id],
                    inline_runs=inline_runs,
                )
            )
            next_span_id += 1
        pages.append(PdfPage(page_number, page_w, page_h, tuple(spans)))

    pages.sort(key=lambda page: page.number)
    if not pages:
        raise RuntimeError("pdftohtml XML contains no pages")
    expected_numbers = list(range(1, len(pages) + 1))
    if [page.number for page in pages] != expected_numbers:
        raise RuntimeError("pdftohtml XML page numbering is not contiguous")
    return PdfDocument(tuple(pages))


def page_number(path: Path) -> int:
    match = re.search(r"(\d+)$", path.stem)
    if match is None:
        raise ValueError(f"Cannot derive page number from {path}")
    return int(match.group(1))


def validate_text_free_pdf(filtered_pdf: Path) -> None:
    result = subprocess.run(
        ["pdftotext", str(filtered_pdf), "-"],
        check=True,
        text=True,
        capture_output=True,
    )
    remaining = result.stdout.strip()
    if remaining:
        excerpt = remaining[:160].replace("\n", " ")
        raise RuntimeError(
            "Ghostscript -dFILTERTEXT background still has extractable text: "
            + excerpt
        )


def render_text_free_backgrounds(expected_pages: int) -> list[Path]:
    filtered_pdf = BUILD_DIR / "text-free-background.pdf"
    subprocess.run(
        [
            "gs",
            "-q",
            "-dSAFER",
            "-dBATCH",
            "-dNOPAUSE",
            "-dFILTERTEXT",
            "-dCompatibilityLevel=1.7",
            "-sDEVICE=pdfwrite",
            f"-sOutputFile={filtered_pdf}",
            str(PDF_PATH),
        ],
        check=True,
    )
    validate_text_free_pdf(filtered_pdf)

    output_pattern = BUILD_DIR / "background-%03d.png"
    subprocess.run(
        [
            "gs",
            "-q",
            "-dSAFER",
            "-dBATCH",
            "-dNOPAUSE",
            "-dFILTERTEXT",
            "-sDEVICE=png16m",
            "-dGraphicsAlphaBits=4",
            "-dTextAlphaBits=4",
            f"-r{RENDER_DPI}",
            f"-sOutputFile={output_pattern}",
            str(filtered_pdf),
        ],
        check=True,
    )
    backgrounds = sorted(
        BUILD_DIR.glob("background-*.png"), key=page_number
    )
    if len(backgrounds) != expected_pages:
        raise RuntimeError(
            f"Ghostscript rendered {len(backgrounds)} backgrounds for "
            f"{expected_pages} PDF pages"
        )
    for background in backgrounds:
        with Image.open(background) as image:
            image.verify()
            if image.width <= 0 or image.height <= 0:
                raise RuntimeError(f"Invalid background image: {background}")
    return backgrounds


def is_msam_triangle(span: TextSpan) -> bool:
    family = strip_subset_prefix(span.font.source_family).upper()
    return family == "MSAM10" and span.text.strip() == "I"


def is_monospace_span(span: TextSpan) -> bool:
    """Return whether a span is source code that must keep hard line breaks."""

    family = strip_subset_prefix(span.font.source_family).lower()
    mapped = span.font.mapped_family.lower()
    return (
        "mono" in family
        or family.startswith("cmtt")
        or "mono" in mapped
        or "courier" in mapped
    )


def line_is_monospace(line: tuple[TextSpan, ...]) -> bool:
    return bool(line) and all(is_monospace_span(span) for span in line)


def vertical_overlap(first: TextSpan, second: TextSpan) -> float:
    return max(0.0, min(first.bottom, second.bottom) - max(first.top, second.top))


def same_visual_line(first: TextSpan, second: TextSpan) -> bool:
    minimum_height = min(first.height, second.height)
    if minimum_height <= 0:
        return False
    overlap_ratio = vertical_overlap(first, second) / minimum_height
    baseline_delta = abs(first.bottom - second.bottom)
    return (
        overlap_ratio >= 0.55
        and baseline_delta <= max(1.5, 0.32 * max(first.height, second.height))
    )


def adjacency_limit(first: TextSpan, second: TextSpan) -> float:
    average_font_size = (first.font.size + second.font.size) / 2.0
    return max(3.0, min(14.0, 0.80 * average_font_size))


def justified_adjacency_limit(first: TextSpan, second: TextSpan) -> float:
    average_font_size = (first.font.size + second.font.size) / 2.0
    # TeX stretches inter-word whitespace when it justifies a narrow table
    # cell.  Apply this wider limit only after a repeated column anchor; using
    # it globally would incorrectly combine independent diagram/footer labels.
    return max(4.0, min(28.0, 2.35 * average_font_size))


def should_group_adjacent(first: TextSpan, second: TextSpan) -> bool:
    if not same_visual_line(first, second):
        return False
    gap = second.left - first.right
    minimum_gap = -max(1.0, 0.12 * min(first.height, second.height))
    return minimum_gap <= gap <= adjacency_limit(first, second)


def build_visual_lines(spans: list[TextSpan]) -> list[list[TextSpan]]:
    visual_lines: list[list[TextSpan]] = []
    for span in sorted(spans, key=lambda item: (item.top, item.left)):
        candidates = [
            line
            for line in visual_lines
            if any(same_visual_line(span, member) for member in line)
        ]
        if not candidates:
            visual_lines.append([span])
            continue
        best = min(
            candidates,
            key=lambda line: min(
                abs(span.bottom - member.bottom) for member in line
            ),
        )
        best.append(span)
    return visual_lines


def repeated_left_anchors(spans: list[TextSpan]) -> tuple[float, ...]:
    """Return x positions that repeatedly start lines/columns on a page."""

    anchors: list[list[TextSpan]] = []
    for span in sorted(spans, key=lambda item: item.left):
        matching = next(
            (
                group
                for group in anchors
                if abs(group[0].left - span.left)
                <= max(1.5, 0.14 * span.font.size)
            ),
            None,
        )
        if matching is None:
            anchors.append([span])
        else:
            matching.append(span)
    return tuple(
        sum(span.left for span in group) / len(group)
        for group in anchors
        if len({round(span.top, 1) for span in group}) >= 2
    )


def starts_at_repeated_anchor(
    span: TextSpan, anchors: tuple[float, ...]
) -> bool:
    return any(
        abs(span.left - anchor) <= max(1.5, 0.14 * span.font.size)
        for anchor in anchors
    )


def is_standard_footer_span(span: TextSpan) -> bool:
    return (span.top / span.page_h) * SLIDE_H_IN >= STANDARD_FOOTER_TOP_IN


def build_line_clusters(spans: list[TextSpan]) -> tuple[list[tuple[TextSpan, ...]], int]:
    """Split each visual line into adjacent styled-span clusters."""

    clusters: list[tuple[TextSpan, ...]] = []
    adjacent_pairs = 0
    anchors = repeated_left_anchors(spans)
    for visual_line in build_visual_lines(spans):
        ordered = sorted(visual_line, key=lambda item: item.left)
        if not ordered:
            continue
        cluster: list[TextSpan] = [ordered[0]]
        for span in ordered[1:]:
            gap = span.left - cluster[-1].right
            minimum_gap = -max(1.0, 0.12 * min(cluster[-1].height, span.height))
            justified_word = (
                same_visual_line(cluster[-1], span)
                and not is_standard_footer_span(cluster[0])
                and starts_at_repeated_anchor(cluster[0], anchors)
                and not starts_at_repeated_anchor(span, anchors)
                and minimum_gap <= gap <= justified_adjacency_limit(cluster[-1], span)
            )
            crosses_column_boundary = (
                starts_at_repeated_anchor(cluster[0], anchors)
                and starts_at_repeated_anchor(span, anchors)
                and span.left - cluster[0].left
                > max(14.0, 1.30 * span.font.size)
            )
            ordinary_word = should_group_adjacent(cluster[-1], span)
            if ordinary_word or (justified_word and not crosses_column_boundary):
                cluster.append(span)
                adjacent_pairs += 1
                continue
            clusters.append(tuple(cluster))
            cluster = [span]
        clusters.append(tuple(cluster))
    clusters.sort(
        key=lambda line: (
            min(span.top for span in line),
            min(span.left for span in line),
        )
    )
    return clusters, adjacent_pairs


def line_left(line: tuple[TextSpan, ...]) -> float:
    return min(span.left for span in line)


def line_right(line: tuple[TextSpan, ...]) -> float:
    return max(span.right for span in line)


def line_top(line: tuple[TextSpan, ...]) -> float:
    return min(span.top for span in line)


def line_bottom(line: tuple[TextSpan, ...]) -> float:
    return max(span.bottom for span in line)


def line_height(line: tuple[TextSpan, ...]) -> float:
    return max(span.height for span in line)


def line_font_size(line: tuple[TextSpan, ...]) -> float:
    return max(span.font.size for span in line)


def line_color(line: tuple[TextSpan, ...]) -> tuple[int, int, int]:
    return max(line, key=lambda span: len(span.text.strip())).font.rgb


def expected_line_text(line: tuple[TextSpan, ...]) -> str:
    parts: list[str] = []
    for index, span in enumerate(line):
        if index:
            parts.append(spacer_text(line[index - 1], span))
        parts.append(span.text)
    return "".join(parts)


def line_starts_with_bullet(
    line: tuple[TextSpan, ...], bullet_spans: list[TextSpan]
) -> bool:
    left = line_left(line)
    size = line_font_size(line)
    for bullet in bullet_spans:
        if not any(same_visual_line(bullet, span) for span in line):
            continue
        gap = left - bullet.right
        if -1.0 <= gap <= max(24.0, 2.4 * size):
            return True
    return False


LIST_MARKER_RE = re.compile(r"(?:\d+|[A-Za-z]|[ivxlcdmIVXLCDM]+)[.)]")
INLINE_LABEL_RE = re.compile(r"[A-Za-z]{1,4}\d+")


def paragraph_lane_left(line: tuple[TextSpan, ...]) -> float:
    """Return the body-column anchor for a labeled paragraph line.

    Description-style entries such as ``RQ1`` contain a short label followed
    by the paragraph body.  PDF extraction reports the first line's left edge
    at the label, while wrapped lines start at the body column.  Treating the
    body column as the paragraph anchor lets the resulting PPTX use one
    wrapping textbox instead of one textbox per visual line.
    """

    ordered = sorted(line, key=lambda span: span.left)
    if len(ordered) >= 2 and INLINE_LABEL_RE.fullmatch(ordered[0].text.strip()):
        return ordered[1].left
    return line_left(line)


def line_starts_with_list_marker(line: tuple[TextSpan, ...]) -> bool:
    """Return true for an explicit standalone list marker such as ``1.``."""

    first = min(line, key=lambda span: span.left)
    return LIST_MARKER_RE.fullmatch(first.text.strip()) is not None


def line_has_separate_sibling(
    line: tuple[TextSpan, ...], all_lines: list[tuple[TextSpan, ...]]
) -> bool:
    """Detect a table/column row that has another non-adjacent cell."""

    for sibling in all_lines:
        if sibling is line:
            continue
        if not any(
            same_visual_line(first, second)
            for first in line
            for second in sibling
        ):
            continue
        if line_right(sibling) < line_left(line):
            gap = line_left(line) - line_right(sibling)
        elif line_right(line) < line_left(sibling):
            gap = line_left(sibling) - line_right(line)
        else:
            continue
        if gap <= max(80.0, 7.0 * max(line_font_size(line), line_font_size(sibling))):
            return True
    return False


def paragraph_join_text(previous: TextSpan, following: TextSpan) -> str:
    previous_text = previous.text.rstrip()
    following_text = following.text.lstrip()
    if not previous_text or not following_text:
        return ""
    if previous_text.endswith(("-", "–", "/")):
        return ""
    if previous.text[-1:].isspace() or following.text[:1].isspace():
        return ""
    return " "


def can_merge_paragraph_line(
    paragraph_lines: list[tuple[TextSpan, ...]],
    following: tuple[TextSpan, ...],
    all_lines: list[tuple[TextSpan, ...]],
    bullet_spans: list[TextSpan],
) -> bool:
    previous = paragraph_lines[-1]
    # Verbatim/code blocks use indentation and explicit newlines as syntax.
    # Keeping each visual code line in its own non-wrapping textbox prevents
    # LibreOffice and Google Slides from replacing those newlines with spaces
    # or reflowing indentation after font substitution.
    if line_is_monospace(previous) or line_is_monospace(following):
        return False
    if line_starts_with_bullet(following, bullet_spans) or line_starts_with_list_marker(
        following
    ):
        return False

    top_delta = line_top(following) - line_top(previous)
    maximum_height = max(line_height(previous), line_height(following))
    if top_delta <= 0.20 * maximum_height or top_delta > 1.75 * maximum_height:
        return False

    previous_size = line_font_size(previous)
    following_size = line_font_size(following)
    if abs(previous_size - following_size) > max(
        1.0, 0.14 * max(previous_size, following_size)
    ):
        return False
    if line_color(previous) != line_color(following):
        return False

    previous_text = expected_line_text(previous).rstrip()
    hyphen_continuation = previous_text.endswith(("-", "–", "/"))
    if line_has_separate_sibling(following, all_lines) and not hyphen_continuation:
        return False

    anchor = paragraph_lines[0]
    anchor_lane_left = paragraph_lane_left(anchor)
    left_tolerance = max(3.0, 0.34 * max(previous_size, following_size))
    left_aligned = abs(line_left(following) - line_left(previous)) <= left_tolerance
    hanging_indent = (
        line_left(following) >= anchor_lane_left
        and line_left(following) - anchor_lane_left
        <= max(18.0, 1.7 * following_size)
    )
    previous_center = (line_left(previous) + line_right(previous)) / 2.0
    following_center = (line_left(following) + line_right(following)) / 2.0
    centered = abs(previous_center - following_center) <= max(
        5.0, 0.55 * following_size
    )
    return left_aligned or hanging_indent or centered


def shares_paragraph_lane(
    paragraph_lines: list[tuple[TextSpan, ...]],
    following: tuple[TextSpan, ...],
) -> bool:
    """Whether two lines occupy the same likely textbox column.

    This deliberately ignores row-start and punctuation rules.  The caller
    first finds the nearest line in the same lane and only then decides if it
    may merge.  That prevents a rejected new table row from attaching to an
    older row farther above.
    """

    previous = paragraph_lines[-1]
    top_delta = line_top(following) - line_top(previous)
    maximum_height = max(line_height(previous), line_height(following))
    if top_delta <= 0.20 * maximum_height or top_delta > 1.75 * maximum_height:
        return False

    previous_size = line_font_size(previous)
    following_size = line_font_size(following)
    if abs(previous_size - following_size) > max(
        1.0, 0.14 * max(previous_size, following_size)
    ):
        return False

    anchor = paragraph_lines[0]
    anchor_lane_left = paragraph_lane_left(anchor)
    left_tolerance = max(3.0, 0.34 * max(previous_size, following_size))
    left_aligned = abs(line_left(following) - line_left(previous)) <= left_tolerance
    hanging_indent = (
        line_left(following) >= anchor_lane_left
        and line_left(following) - anchor_lane_left
        <= max(18.0, 1.7 * following_size)
    )
    previous_center = (line_left(previous) + line_right(previous)) / 2.0
    following_center = (line_left(following) + line_right(following)) / 2.0
    centered = abs(previous_center - following_center) <= max(
        5.0, 0.55 * following_size
    )
    return left_aligned or hanging_indent or centered


def merge_line_clusters(
    lines: list[tuple[TextSpan, ...]], bullet_spans: list[TextSpan]
) -> list[list[tuple[TextSpan, ...]]]:
    """Merge PDF visual lines into paragraph-sized editable blocks."""

    paragraphs: list[list[tuple[TextSpan, ...]]] = []
    for line in lines:
        lane_candidates = [
            paragraph
            for paragraph in paragraphs
            if shares_paragraph_lane(paragraph, line)
        ]
        if not lane_candidates:
            paragraphs.append([line])
            continue
        nearest = min(
            lane_candidates,
            key=lambda paragraph: line_top(line) - line_top(paragraph[-1]),
        )
        if can_merge_paragraph_line(nearest, line, lines, bullet_spans):
            nearest.append(line)
        else:
            paragraphs.append([line])
    paragraphs.sort(
        key=lambda paragraph: (
            line_top(paragraph[0]),
            line_left(paragraph[0]),
        )
    )
    return paragraphs


def infer_text_alignment(
    page: PdfPage,
    paragraph_lines: list[tuple[TextSpan, ...]],
    bullet_spans: list[TextSpan],
) -> str:
    """Infer semantic paragraph alignment from the PDF line geometry."""

    if not paragraph_lines:
        return "left"
    first = paragraph_lines[0]
    if line_starts_with_bullet(first, bullet_spans) or line_starts_with_list_marker(first):
        return "left"
    if all(is_standard_footer_span(span) for line in paragraph_lines for span in line):
        return "left"

    # The title page is a deliberately centered composition.  Unlike normal
    # frame titles it has no left-aligned header/footer structure from which a
    # single-line alignment could be inferred.
    if page.number == 1:
        return "center"

    sizes = [line_font_size(line) for line in paragraph_lines]
    tolerance = max(2.0, 0.22 * max(sizes))
    centers = [(line_left(line) + line_right(line)) / 2.0 for line in paragraph_lines]
    lefts = [line_left(line) for line in paragraph_lines]
    if len(paragraph_lines) >= 2:
        center_spread = max(centers) - min(centers)
        left_spread = max(lefts) - min(lefts)
        if center_spread <= tolerance and left_spread > tolerance:
            return "center"
        return "left"

    line_width = line_right(first) - line_left(first)
    near_page_center = abs(centers[0] - page.width / 2.0) <= tolerance
    if (
        near_page_center
        and line_width <= 0.65 * page.width
        and line_top(first) / page.height * SLIDE_H_IN >= 1.0
    ):
        return "center"
    return "left"


def plan_conversion(document: PdfDocument) -> ConversionPlan:
    groups_by_page: dict[int, tuple[TextGroup, ...]] = {}
    bullets_by_page: dict[int, tuple[BulletMarker, ...]] = {}
    assignments: dict[int, str] = {}
    assignment_events: list[int] = []
    next_group_id = 1
    adjacent_pairs = 0
    multi_span_groups = 0
    svs_citation_groups = 0

    for page in document.pages:
        bullet_spans = [span for span in page.spans if is_msam_triangle(span)]
        text_spans = [span for span in page.spans if not is_msam_triangle(span)]
        bullets = tuple(BulletMarker(page.number, span) for span in bullet_spans)
        for bullet in bullets:
            assignments[bullet.span.span_id] = f"bullet-{bullet.span.span_id}"
            assignment_events.append(bullet.span.span_id)

        line_clusters, page_adjacent_pairs = build_line_clusters(text_spans)
        adjacent_pairs += page_adjacent_pairs
        paragraph_blocks = merge_line_clusters(line_clusters, bullet_spans)

        groups: list[TextGroup] = []
        for paragraph_lines in paragraph_blocks:
            ordered_spans = tuple(
                span
                for line in paragraph_lines
                for span in sorted(line, key=lambda item: item.left)
            )
            line_starts = tuple(
                min(line, key=lambda item: item.left).span_id
                for line in paragraph_lines[1:]
            )
            group = TextGroup(
                next_group_id,
                page.number,
                ordered_spans,
                line_starts,
                infer_text_alignment(page, paragraph_lines, bullet_spans),
            )
            groups.append(group)
            adjacent_pairs += max(0, len(paragraph_lines) - 1)
            next_group_id += 1

        groups.sort(key=lambda group: (group.spans[0].top, group.spans[0].left))
        for group in groups:
            assignment_name = f"text-group-{group.group_id}"
            for span in group.spans:
                if span.span_id in assignments:
                    raise RuntimeError(
                        f"Text span {span.span_id} was assigned more than once"
                    )
                assignments[span.span_id] = assignment_name
                assignment_events.append(span.span_id)
            if len(group.spans) > 1:
                multi_span_groups += 1
            texts = [span.text.strip() for span in group.spans]
            if any("SVS" in text for text in texts) and any(
                text.startswith("[3]") for text in texts
            ):
                svs_citation_groups += 1

        groups_by_page[page.number] = tuple(groups)
        bullets_by_page[page.number] = bullets

    counts = Counter(assignment_events)
    source_ids = {span.span_id for span in document.spans}
    assigned_ids = set(counts)
    missing = sorted(source_ids - assigned_ids)
    unexpected = sorted(assigned_ids - source_ids)
    duplicated = sorted(span_id for span_id, count in counts.items() if count != 1)
    if missing or unexpected or duplicated:
        raise RuntimeError(
            "Invalid source-span assignment: "
            f"missing={missing[:10]}, unexpected={unexpected[:10]}, "
            f"duplicated={duplicated[:10]}"
        )

    return ConversionPlan(
        groups_by_page=groups_by_page,
        bullets_by_page=bullets_by_page,
        assignments=assignments,
        adjacent_pairs=adjacent_pairs,
        multi_span_groups=multi_span_groups,
        svs_citation_groups=svs_citation_groups,
    )


def resolve_visual_regions(
    document: PdfDocument,
) -> dict[int, tuple[VisualRegion, ...]]:
    resolved: dict[int, tuple[VisualRegion, ...]] = {}
    region_names: set[str] = set()

    for spec in SLIDE_VISUAL_REGION_SPECS:
        matched_pages = [
            page
            for page in document.pages
            if all(
                required in " ".join(span.text for span in page.spans)
                for required in spec.required_text
            )
        ]
        if len(matched_pages) > 1:
            raise RuntimeError(
                "Visual-region slide match is ambiguous for "
                f"{spec.required_text!r}: pages "
                f"{[page.number for page in matched_pages]}"
            )
        if not matched_pages:
            continue

        page = matched_pages[0]
        for region in spec.regions:
            if region.name in region_names:
                raise RuntimeError(
                    f"Duplicate visual-region shape name: {region.name}"
                )
            if (
                region.left < 0
                or region.top < 0
                or region.width <= 0
                or region.height <= 0
                or region.left + region.width > SLIDE_W_IN
                or region.top + region.height > SLIDE_H_IN
            ):
                raise RuntimeError(
                    f"Visual region is outside the slide: {region!r}"
                )
            region_names.add(region.name)
        if page.number in resolved:
            raise RuntimeError(
                f"Multiple visual-region specs matched page {page.number}"
            )
        resolved[page.number] = spec.regions

    return resolved


def x_inches(value: float, page: PdfPage) -> float:
    return value / page.width * SLIDE_W_IN


def y_inches(value: float, page: PdfPage) -> float:
    return value / page.height * SLIDE_H_IN


def font_size_points(span: TextSpan) -> float:
    return span.font.size / span.page_w * SLIDE_W_IN * 72.0


def remove_shape_line(shape) -> None:
    line = shape._element.spPr.get_or_add_ln()
    for child in list(line):
        line.remove(child)
    line.append(OxmlElement("a:noFill"))


def add_powerpoint_run(paragraph, span: TextSpan, inline: InlineTextRun) -> None:
    run = paragraph.add_run()
    run.text = inline.text
    run.font.name = span.font.mapped_family
    run.font.size = Pt(font_size_points(span))
    run.font.bold = span.font.family_bold or inline.bold
    run.font.italic = span.font.family_italic or inline.italic
    run.font.color.rgb = RGBColor(*span.font.rgb)


def spacer_text(first: TextSpan, second: TextSpan) -> str:
    if first.text[-1:].isspace() or second.text[:1].isspace():
        return ""
    gap = second.left - first.right
    if gap <= max(0.4, 0.04 * first.font.size):
        return ""
    # A logical paragraph must use semantic whitespace.  Reproducing the PDF's
    # justified x-gap with several literal spaces makes Google Slides and
    # LibreOffice wrap at different positions after font substitution.
    return " "


def add_spacer_run(paragraph, first: TextSpan, second: TextSpan) -> int:
    text = spacer_text(first, second)
    if not text:
        return 0
    run = paragraph.add_run()
    run.text = text
    run.font.name = first.font.mapped_family
    run.font.size = Pt(font_size_points(first))
    run.font.bold = first.font.family_bold
    run.font.italic = first.font.family_italic
    run.font.color.rgb = RGBColor(*first.font.rgb)
    return 1


def add_paragraph_join_run(paragraph, first: TextSpan, second: TextSpan) -> int:
    text = paragraph_join_text(first, second)
    if not text:
        return 0
    run = paragraph.add_run()
    run.text = text
    run.font.name = first.font.mapped_family
    run.font.size = Pt(font_size_points(first))
    run.font.bold = first.font.family_bold
    run.font.italic = first.font.family_italic
    run.font.color.rgb = RGBColor(*first.font.rgb)
    return 1


def expected_group_text(group: TextGroup) -> str:
    parts: list[str] = []
    line_starts = set(group.line_start_span_ids)
    preserve_centered_lines = has_centered_visual_line_geometry(group)
    for index, span in enumerate(group.spans):
        if index:
            previous = group.spans[index - 1]
            parts.append(
                (
                    "\v"
                    if preserve_centered_lines
                    else paragraph_join_text(previous, span)
                )
                if span.span_id in line_starts
                else spacer_text(previous, span)
            )
        parts.append(span.text)
    return "".join(parts)


def has_centered_visual_line_geometry(group: TextGroup) -> bool:
    """Whether a centered group contains intentionally centered PDF lines.

    Center alignment alone is insufficient: the title-page fallback and some
    narrow table cells can be centered even when their PDF lines are ordinary
    wrapped prose.  Preserve hard line breaks only when the source-line
    bounding boxes themselves share a common center.
    """

    if group.alignment != "center" or not group.line_start_span_ids:
        return False

    line_starts = set(group.line_start_span_ids)
    lines: list[list[TextSpan]] = [[]]
    for index, span in enumerate(group.spans):
        if index and span.span_id in line_starts:
            lines.append([])
        lines[-1].append(span)

    centers = [
        (min(span.left for span in line) + max(span.right for span in line)) / 2.0
        for line in lines
    ]
    largest_font = max(font_size_points(span) for span in group.spans)
    return max(centers) - min(centers) <= max(3.0, 0.30 * largest_font)


def is_standard_footer_group(page: PdfPage, group: TextGroup) -> bool:
    source_top = min(span.top for span in group.spans)
    return y_inches(source_top, page) >= STANDARD_FOOTER_TOP_IN


def add_text_group(slide, page: PdfPage, group: TextGroup) -> int:
    left = min(span.left for span in group.spans)
    top = min(span.top for span in group.spans)
    right = max(span.right for span in group.spans)
    bottom = max(span.bottom for span in group.spans)
    x = x_inches(left, page)
    y = y_inches(top, page)
    source_width = x_inches(right - left, page)
    source_height = y_inches(bottom - top, page)
    largest_font_height = max(font_size_points(span) for span in group.spans) / 72.0
    is_multiline = bool(group.line_start_span_ids)
    is_monospace = all(is_monospace_span(span) for span in group.spans)
    width_padding = (
        0.14
        if is_multiline
        else max(0.12, 0.60 * largest_font_height, 0.08 * source_width)
    )
    if is_monospace:
        width_padding = max(width_padding, 0.22)
    height_padding = largest_font_height * (0.52 if is_multiline else 0.30)
    if group.alignment == "center":
        x = max(0.0, x - width_padding / 2.0)
    width = min(SLIDE_W_IN - x, max(0.03, source_width + width_padding))
    is_frame_title = y < 1.0 and max(
        font_size_points(span) for span in group.spans
    ) >= 20.0
    title_left_inset = 0.0
    if is_frame_title:
        # A title's PDF glyph bounds are not its layout bounds.  Give it the
        # full content width so substituted fonts do not wrap a nominally
        # single-line title above the slide in Google Slides/LibreOffice.
        if group.alignment == "center":
            x = 0.30
            width = SLIDE_W_IN - 0.60
        else:
            # Give the glyphs a small safety area to the left of the visible
            # title anchor. LibreOffice can otherwise clip the negative left
            # bearing of the first glyph after import/save (notably the "P" in
            # "Preliminary"). The matching text-frame inset keeps the visible
            # text at the source PDF position.
            title_left_inset = min(0.10, x)
            x -= title_left_inset
            # Use nearly all remaining slide width. LibreOffice can otherwise
            # rewrap a title after save even when the first import rendered it
            # on one line.
            width = max(width, SLIDE_W_IN - x - 0.05)
    desired_height = max(
        0.04,
        source_height + height_padding,
        largest_font_height * (1.42 if is_multiline else 1.28),
    )
    if is_standard_footer_group(page, group):
        height = min(
            SLIDE_H_IN - FOOTER_BOTTOM_MARGIN_IN, desired_height
        )
        y = min(
            y,
            SLIDE_H_IN - FOOTER_BOTTOM_MARGIN_IN - height,
        )
    else:
        height = min(SLIDE_H_IN - y, desired_height)
    textbox = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    textbox.name = f"PDF_TEXT_GROUP_{group.group_id}"
    textbox.fill.background()
    remove_shape_line(textbox)
    frame = textbox.text_frame
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    # Frame titles are semantically one line.  Explicitly disabling wrapping
    # survives the Google Slides/LibreOffice import-save path better than
    # relying on a font-specific glyph-width fit. Body paragraphs still wrap.
    frame.word_wrap = not (is_frame_title or is_monospace)
    frame.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    frame.margin_left = Inches(title_left_inset)
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = (
        PP_ALIGN.CENTER if group.alignment == "center" else PP_ALIGN.LEFT
    )
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    if line_starts_with_list_marker((group.spans[0],)) and len(group.spans) >= 2:
        marker = group.spans[0]
        body = group.spans[1]
        if body.span_id not in set(group.line_start_span_ids):
            hanging_indent = x_inches(max(0.0, body.left - marker.left), page)
            paragraph_properties = paragraph._p.get_or_add_pPr()
            hanging_indent_emu = int(Inches(hanging_indent))
            paragraph_properties.set("marL", str(hanging_indent_emu))
            paragraph_properties.set("indent", str(-hanging_indent_emu))

    run_count = 0
    line_starts = set(group.line_start_span_ids)
    preserve_centered_lines = has_centered_visual_line_geometry(group)
    for index, span in enumerate(group.spans):
        prefix = ""
        if index:
            previous = group.spans[index - 1]
            if span.span_id in line_starts and preserve_centered_lines:
                # Centered diagram labels and cover-page blocks use their PDF
                # line boundaries as intentional layout, not prose wrapping.
                paragraph.add_line_break()
            else:
                prefix = (
                    paragraph_join_text(previous, span)
                    if span.span_id in line_starts
                    else spacer_text(previous, span)
                )
        for inline_index, inline in enumerate(span.inline_runs):
            # LibreOffice and Google Slides may discard a whitespace-only run
            # between two styled spans.  Prefix semantic whitespace to the
            # following visible run so the saved PPTX retains word boundaries.
            visible_inline = (
                InlineTextRun(prefix + inline.text, inline.bold, inline.italic)
                if prefix and inline_index == 0
                else inline
            )
            add_powerpoint_run(paragraph, span, visible_inline)
            run_count += 1
    return run_count


def add_native_triangle(slide, page: PdfPage, marker: BulletMarker) -> None:
    span = marker.span
    box_x = x_inches(span.left, page)
    box_y = y_inches(span.top, page)
    box_w = x_inches(span.width, page)
    box_h = y_inches(span.height, page)
    side = max(0.035, min(box_w * 0.72, box_h * 0.56))
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        Inches(box_x + (box_w - side) / 2.0),
        Inches(box_y + (box_h - side) / 2.0),
        Inches(side),
        Inches(side),
    )
    shape.name = f"PDF_BULLET_SPAN_{span.span_id}"
    shape.rotation = 90
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*span.font.rgb)
    shape.line.fill.background()


def write_manifest(
    document: PdfDocument,
    plan: ConversionPlan,
    visual_regions_by_page: dict[int, tuple[VisualRegion, ...]],
) -> Path:
    manifest_path = BUILD_DIR / MANIFEST_NAME
    payload = {
        "source_pdf": str(PDF_PATH),
        "background_filter": "Ghostscript -dFILTERTEXT",
        "validation": {
            "source_span_count": len(document.spans),
            "assigned_once_count": len(plan.assignments),
            "adjacent_pairs_grouped": plan.adjacent_pairs,
            "multi_span_textboxes": plan.multi_span_groups,
            "svs_citation_textboxes": plan.svs_citation_groups,
            "background_extractable_text_characters": 0,
            "independent_visual_region_count": sum(
                len(regions) for regions in visual_regions_by_page.values()
            ),
        },
        "pages": [
            {
                "page": page.number,
                "width": page.width,
                "height": page.height,
                "visual_regions": [
                    {
                        "name": region.name,
                        "left_inches": region.left,
                        "top_inches": region.top,
                        "width_inches": region.width,
                        "height_inches": region.height,
                    }
                    for region in visual_regions_by_page.get(page.number, ())
                ],
                "text_groups": [
                    {
                        "group_id": group.group_id,
                        "alignment": group.alignment,
                        "line_start_span_ids": list(group.line_start_span_ids),
                        "text": expected_group_text(group),
                    }
                    for group in plan.groups_by_page.get(page.number, ())
                ],
                "spans": [
                    {
                        "span_id": span.span_id,
                        "assignment": plan.assignments[span.span_id],
                        "text": span.text,
                        "left": span.left,
                        "top": span.top,
                        "width": span.width,
                        "height": span.height,
                        "font_size": span.font.size,
                        "source_family": span.font.source_family,
                        "mapped_family": span.font.mapped_family,
                        "rgb": list(span.font.rgb),
                        "family_bold": span.font.family_bold,
                        "family_italic": span.font.family_italic,
                        "inline_runs": [
                            {
                                "text": inline.text,
                                "bold": inline.bold,
                                "italic": inline.italic,
                            }
                            for inline in span.inline_runs
                        ],
                    }
                    for span in page.spans
                ],
            }
            for page in document.pages
        ],
    }
    manifest_path.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    return manifest_path


def visual_region_pixel_box(
    region: VisualRegion, image_width: int, image_height: int
) -> tuple[int, int, int, int]:
    left = round(region.left / SLIDE_W_IN * image_width)
    top = round(region.top / SLIDE_H_IN * image_height)
    right = round(
        (region.left + region.width) / SLIDE_W_IN * image_width
    )
    bottom = round(
        (region.top + region.height) / SLIDE_H_IN * image_height
    )
    if not (0 <= left < right <= image_width):
        raise RuntimeError(
            f"Invalid horizontal crop for {region.name}: "
            f"{(left, right)} of {image_width}"
        )
    if not (0 <= top < bottom <= image_height):
        raise RuntimeError(
            f"Invalid vertical crop for {region.name}: "
            f"{(top, bottom)} of {image_height}"
        )
    return left, top, right, bottom


def prepare_visual_region_images(
    background: Path,
    page: PdfPage,
    regions: tuple[VisualRegion, ...],
) -> tuple[Path, tuple[PreparedVisualRegion, ...]]:
    if not regions:
        return background, ()

    with Image.open(background) as image:
        source = image.convert("RGB")

    base = source.copy()
    prepared: list[PreparedVisualRegion] = []
    pixel_boxes: list[tuple[int, int, int, int]] = []
    for region in regions:
        box = visual_region_pixel_box(region, source.width, source.height)
        for previous in pixel_boxes:
            horizontal_overlap = max(box[0], previous[0]) < min(
                box[2], previous[2]
            )
            vertical_overlap = max(box[1], previous[1]) < min(
                box[3], previous[3]
            )
            if horizontal_overlap and vertical_overlap:
                raise RuntimeError(
                    f"Overlapping visual-region crops on page {page.number}: "
                    f"{previous} and {box}"
                )
        pixel_boxes.append(box)

        crop_path = BUILD_DIR / (
            f"visual-region-{page.number:03d}-"
            f"{region.name.lower()}.png"
        )
        cropped = source.crop(box)
        if all(extrema == (255, 255) for extrema in cropped.getextrema()):
            raise RuntimeError(
                f"Visual-region crop is blank on page {page.number}: "
                f"{region.name}"
            )
        cropped.save(crop_path)
        base.paste((255, 255, 255), box)
        left, top, right, bottom = box
        prepared.append(
            PreparedVisualRegion(
                region=region,
                image_path=crop_path,
                left=left / source.width * SLIDE_W_IN,
                top=top / source.height * SLIDE_H_IN,
                width=(right - left) / source.width * SLIDE_W_IN,
                height=(bottom - top) / source.height * SLIDE_H_IN,
            )
        )

    base_path = BUILD_DIR / f"ppt-background-{page.number:03d}.png"
    base.save(base_path)
    return base_path, tuple(prepared)


def build_pptx(
    document: PdfDocument,
    backgrounds: list[Path],
    plan: ConversionPlan,
    visual_regions_by_page: dict[int, tuple[VisualRegion, ...]],
    output_path: Path,
) -> BuildStats:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    blank = prs.slide_layouts[6]

    textboxes = 0
    text_runs = 0
    bullet_shapes = 0
    visual_region_images = 0
    for page, background in zip(document.pages, backgrounds):
        slide = prs.slides.add_slide(blank)
        ppt_background, prepared_regions = prepare_visual_region_images(
            background,
            page,
            visual_regions_by_page.get(page.number, ()),
        )
        picture = slide.shapes.add_picture(
            str(ppt_background),
            0,
            0,
            width=prs.slide_width,
            height=prs.slide_height,
        )
        picture.name = f"TEXT_FREE_BACKGROUND_{page.number}"

        for prepared in prepared_regions:
            region_picture = slide.shapes.add_picture(
                str(prepared.image_path),
                Inches(prepared.left),
                Inches(prepared.top),
                width=Inches(prepared.width),
                height=Inches(prepared.height),
            )
            region_picture.name = prepared.region.name
            visual_region_images += 1

        for group in plan.groups_by_page.get(page.number, ()):
            text_runs += add_text_group(slide, page, group)
            textboxes += 1
        for marker in plan.bullets_by_page.get(page.number, ()):
            add_native_triangle(slide, page, marker)
            bullet_shapes += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return BuildStats(
        textboxes=textboxes,
        text_runs=text_runs,
        bullet_shapes=bullet_shapes,
        background_images=len(backgrounds),
        visual_region_images=visual_region_images,
    )


def validate_saved_pptx(
    document: PdfDocument,
    plan: ConversionPlan,
    stats: BuildStats,
    visual_regions_by_page: dict[int, tuple[VisualRegion, ...]],
    output_path: Path,
) -> None:
    prs = Presentation(output_path)
    if len(prs.slides) != len(document.pages):
        raise RuntimeError(
            f"Saved PPTX has {len(prs.slides)} slides; expected {len(document.pages)}"
        )

    background_count = 0
    visual_region_count = 0
    text_group_count = 0
    bullet_count = 0
    saved_group_ids: list[int] = []
    saved_bullet_ids: list[int] = []
    expected_groups = {
        group.group_id: group
        for groups in plan.groups_by_page.values()
        for group in groups
    }
    footer_group_ids = {
        group.group_id
        for page in document.pages
        for group in plan.groups_by_page.get(page.number, ())
        if is_standard_footer_group(page, group)
    }
    for page, slide in zip(document.pages, prs.slides):
        pictures = [
            shape
            for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ]
        backgrounds = [
            shape
            for shape in pictures
            if shape.name == f"TEXT_FREE_BACKGROUND_{page.number}"
        ]
        if len(backgrounds) != 1:
            raise RuntimeError(
                f"Slide {page.number} has {len(backgrounds)} background pictures"
            )
        if backgrounds[0].image.ext.lower() != "png":
            raise RuntimeError(
                f"Slide {page.number} background is not a raster PNG"
            )
        background_count += 1

        expected_regions = {
            region.name: region
            for region in visual_regions_by_page.get(page.number, ())
        }
        region_pictures = [
            shape for shape in pictures if shape.name in expected_regions
        ]
        unexpected_pictures = [
            shape
            for shape in pictures
            if (
                shape.name != f"TEXT_FREE_BACKGROUND_{page.number}"
                and shape.name not in expected_regions
            )
        ]
        if unexpected_pictures:
            raise RuntimeError(
                f"Slide {page.number} has unexpected picture shapes: "
                f"{[shape.name for shape in unexpected_pictures]}"
            )
        if Counter(shape.name for shape in region_pictures) != Counter(
            expected_regions.keys()
        ):
            raise RuntimeError(
                f"Slide {page.number} visual regions are missing or duplicated"
            )
        for shape in region_pictures:
            if shape.image.ext.lower() != "png":
                raise RuntimeError(
                    f"Visual region {shape.name} is not a raster PNG"
                )
            expected = expected_regions[shape.name]
            actual_geometry = (
                shape.left / Inches(1),
                shape.top / Inches(1),
                shape.width / Inches(1),
                shape.height / Inches(1),
            )
            expected_geometry = (
                expected.left,
                expected.top,
                expected.width,
                expected.height,
            )
            if any(
                abs(actual - target) > 0.012
                for actual, target in zip(
                    actual_geometry, expected_geometry
                )
            ):
                raise RuntimeError(
                    f"Visual region {shape.name} geometry changed: "
                    f"expected {expected_geometry}, found {actual_geometry}"
                )
            visual_region_count += 1

        for shape in slide.shapes:
            group_match = re.fullmatch(r"PDF_TEXT_GROUP_(\d+)", shape.name)
            bullet_match = re.fullmatch(r"PDF_BULLET_SPAN_(\d+)", shape.name)
            if group_match:
                text_group_count += 1
                group_id = int(group_match.group(1))
                saved_group_ids.append(group_id)
                if not shape.has_text_frame:
                    raise RuntimeError(
                        f"Saved text group {group_id} has no text frame"
                    )
                expected_group = expected_groups[group_id]
                expected_is_frame_title = (
                    y_inches(min(span.top for span in expected_group.spans), page) < 1.0
                    and max(font_size_points(span) for span in expected_group.spans) >= 20.0
                )
                expected_is_monospace = all(
                    is_monospace_span(span) for span in expected_group.spans
                )
                expected_word_wrap = not (
                    expected_is_frame_title or expected_is_monospace
                )
                if shape.text_frame.word_wrap is not expected_word_wrap:
                    raise RuntimeError(
                        f"Saved text group {group_id} has incorrect word-wrap state"
                    )
                expected_text = expected_group_text(expected_group)
                if shape.text != expected_text:
                    raise RuntimeError(
                        f"Saved text group {group_id} content changed: "
                        f"expected {expected_text!r}, found {shape.text!r}"
                    )
                expected_alignment = (
                    PP_ALIGN.CENTER
                    if expected_groups[group_id].alignment == "center"
                    else PP_ALIGN.LEFT
                )
                if shape.text_frame.paragraphs[0].alignment != expected_alignment:
                    raise RuntimeError(
                        f"Saved text group {group_id} alignment changed"
                    )
                if group_id in footer_group_ids:
                    bottom_inches = (shape.top + shape.height) / Inches(1)
                    if bottom_inches > (
                        SLIDE_H_IN - FOOTER_BOTTOM_MARGIN_IN + 0.002
                    ):
                        raise RuntimeError(
                            f"Saved footer group {group_id} reaches "
                            f"{bottom_inches:.3f} inches"
                        )
            if bullet_match:
                bullet_count += 1
                saved_bullet_ids.append(int(bullet_match.group(1)))

    if background_count != stats.background_images:
        raise RuntimeError("Saved PPTX background count changed")
    if visual_region_count != stats.visual_region_images:
        raise RuntimeError("Saved PPTX visual-region count changed")
    if text_group_count != stats.textboxes:
        raise RuntimeError("Saved PPTX textbox count changed")
    if bullet_count != stats.bullet_shapes:
        raise RuntimeError("Saved PPTX bullet count changed")
    expected_group_ids = set(expected_groups)
    expected_bullet_ids = {
        marker.span.span_id
        for markers in plan.bullets_by_page.values()
        for marker in markers
    }
    if Counter(saved_group_ids) != Counter(expected_group_ids):
        raise RuntimeError("Saved PPTX text-group IDs are missing or duplicated")
    if Counter(saved_bullet_ids) != Counter(expected_bullet_ids):
        raise RuntimeError("Saved PPTX bullet-span IDs are missing or duplicated")
    if len(plan.assignments) != len(document.spans):
        raise RuntimeError("Saved PPTX validation found unassigned source spans")


def resolve_slide_path(path: Path) -> Path:
    return path if path.is_absolute() else SLIDES_DIR / path


def validate_distinct_paths(pdf_path: Path, output_path: Path,
                            notes_path: Path) -> None:
    resolved = {
        "PDF input": pdf_path.resolve(),
        "PPTX output": output_path.resolve(),
        "speaker notes": notes_path.resolve(),
    }
    labels = list(resolved)
    for index, first_label in enumerate(labels):
        for second_label in labels[index + 1:]:
            same_inode = False
            try:
                same_inode = os.path.samefile(
                    resolved[first_label], resolved[second_label]
                )
            except FileNotFoundError:
                pass
            if (
                resolved[first_label] == resolved[second_label]
                or same_inode
            ):
                raise ValueError(
                    f"{first_label} and {second_label} must use distinct paths: "
                    f"{resolved[first_label]} and {resolved[second_label]}"
                )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate a PDF-matched, all-editable-text PowerPoint deck."
    )
    parser.add_argument("--pdf", type=Path, default=Path("main.pdf"))
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("NDNSF_proposal_hybrid_editable.pptx"),
    )
    parser.add_argument("--notes-tex", type=Path, default=Path("speaker_notes.tex"))
    parser.add_argument(
        "--no-notes",
        action="store_true",
        help="Generate the deck without injecting speaker notes.",
    )
    parser.add_argument(
        "--build-dir",
        type=Path,
        default=Path("build/hybrid_editable_pptx"),
    )
    return parser.parse_args()


def main() -> None:
    global PDF_PATH, OUTPUT_PATH, BUILD_DIR
    args = parse_args()
    PDF_PATH = resolve_slide_path(args.pdf)
    OUTPUT_PATH = resolve_slide_path(args.output)
    BUILD_DIR = resolve_slide_path(args.build_dir)
    notes_path = resolve_slide_path(args.notes_tex)

    require_external_tools()
    validate_distinct_paths(PDF_PATH, OUTPUT_PATH, notes_path)
    reset_build_dir((PDF_PATH, OUTPUT_PATH, notes_path))
    document = extract_styled_text()
    plan = plan_conversion(document)
    visual_regions_by_page = resolve_visual_regions(document)
    backgrounds = render_text_free_backgrounds(len(document.pages))
    manifest_path = write_manifest(document, plan, visual_regions_by_page)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    descriptor, staged_name = tempfile.mkstemp(
        prefix=f".{OUTPUT_PATH.name}.",
        suffix=".tmp",
        dir=str(OUTPUT_PATH.parent),
    )
    os.close(descriptor)
    staged_output = Path(staged_name)
    output_mode = 0o644
    try:
        stats = build_pptx(
            document,
            backgrounds,
            plan,
            visual_regions_by_page,
            staged_output,
        )
        if not args.no_notes:
            inject_notes(staged_output, notes=load_notes_from_tex(notes_path))
        validate_saved_pptx(
            document,
            plan,
            stats,
            visual_regions_by_page,
            staged_output,
        )
        staged_output.chmod(output_mode)
        staged_output.replace(OUTPUT_PATH)
    finally:
        if staged_output.exists():
            staged_output.unlink()

    print(
        "Validation PASS: "
        f"{len(plan.assignments)}/{len(document.spans)} source spans assigned "
        "exactly once; 0 extractable text characters in FILTERTEXT background"
    )
    print(
        f"Created {stats.textboxes} textboxes with {stats.text_runs} runs, "
        f"{stats.bullet_shapes} native triangle bullets, "
        f"{stats.visual_region_images} independent visual regions, "
        f"{plan.multi_span_groups} multi-span groups, and "
        f"{plan.svs_citation_groups} SVS citation groups"
    )
    print(f"Manifest: {manifest_path}")
    print(f"Wrote {OUTPUT_PATH} from {len(document.pages)} pages")


if __name__ == "__main__":
    main()
